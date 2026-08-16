#!/usr/bin/env python3
"""
Convert EirGrid's half-hourly wind dispatch-down files into the compact
monthly series the site ships.

    python3 tools/dd_convert.py DD-HH-*.xlsx

Download the files from the "DD Half-Hourly Data" section of
https://www.eirgrid.ie/grid/system-and-renewable-data-reports - they sit
behind a JavaScript accordion, so they cannot be fetched by URL pattern:
the version suffix changes without notice (V7, v10) and a guessed URL
would rot silently rather than fail loudly. Closed years never change;
re-run this only when the current year's file is refreshed.

Wind only. Solar coverage does not start until 2023 and solar is about a
tenth of the volume, so including it would give a series whose
denominator changes shape midway.
"""
import collections
import json
import sys

import openpyxl

COLS = {"avail": 4, "output": 5, "hifrq": 6, "rocof": 7, "snsp": 8,
        "trans": 9, "test": 10, "dd": 11, "curt": 12, "cons": 13,
        "other": 14}
OUT = "docs/dispatch_down_monthly.json"


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def main(paths):
    agg = collections.defaultdict(lambda: collections.defaultdict(float))
    for path in sorted(paths):
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        header = list(next(rows))
        # The schema has held across every year and version to date. If
        # it ever moves, fail here rather than write silent nonsense.
        assert header[0] == "UT_TYPE" and header[2] == "HH_TIMESTAMP" \
            and header[9] == "Sum of TRANS_CONSTR_MWH", (path, header)
        n = 0
        for r in rows:
            if not r[0] or r[0] != "Wind":
                continue
            key = (str(r[2])[:7], r[1])
            for name, i in COLS.items():
                agg[key][name] += num(r[i])
            n += 1
        wb.close()
        print(f"{path}: {n} wind rows")
    months = sorted({k[0] for k in agg})
    out = {"months": months, "unit": "GWh", "technology": "Wind",
           "jurisdictions": {}}
    for j in ("IE", "NI"):
        out["jurisdictions"][j] = {
            name: [round(agg[(m, j)][name] / 1000.0, 2) for m in months]
            for name in COLS}
    with open(OUT, "w") as fh:
        json.dump(out, fh, separators=(",", ":"))
    print(f"wrote {OUT}: {len(months)} months, {months[0]}..{months[-1]}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    main(sys.argv[1:])
